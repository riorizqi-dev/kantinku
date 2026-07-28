from pptx import Presentation
from pathlib import Path

p = Presentation(r"docs\KantinKu-Alur-Sistem-ERD.pptx")
print("slides", len(p.slides))
print("size", Path(r"docs\KantinKu-Alur-Sistem-ERD.pptx").stat().st_size)
for i, s in enumerate(p.slides, 1):
    texts = []
    imgs = 0
    for sh in s.shapes:
        if hasattr(sh, "text") and sh.text and sh.text.strip():
            texts.append(sh.text.strip()[:100].replace("\n", " / "))
        if sh.shape_type is not None and "PICTURE" in str(sh.shape_type):
            imgs += 1
        # shape_type 13 is picture
        try:
            if int(sh.shape_type) == 13:
                imgs += 1
        except Exception:
            pass
    title = texts[0] if texts else "(no text)"
    print(f"{i}: imgs~{imgs} | {title}")
